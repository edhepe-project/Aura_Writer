"""
Aura Writer — DOCX Exporter Module
Generación de borradores en formato Microsoft Word (.docx) con estilo editorial.
"""

import os
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Mm, Inches, RGBColor
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH

from tools.exporters.base_exporter import clean_html


class DOCXExporter:
    """Motor de exportación de borradores en formato DOCX."""

    def __init__(self, meta: dict, temp_dir: str):
        self.meta = meta
        self.temp_dir = temp_dir

    def export(self, project_data: dict, output_path: str):
        """Genera el documento Word (.docx) con estilo manuscrito."""
        doc = Document()
        section = doc.sections[0]
        section.page_width = Mm(148)
        section.page_height = Mm(210)
        section.top_margin = Mm(25)
        section.bottom_margin = Mm(25)
        section.left_margin = Mm(25)
        section.right_margin = Mm(20)

        style = doc.styles["Normal"]
        style.font.name = "Courier New"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(44, 36, 22)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

        title = self.meta.get("title", "Borrador")
        author = self.meta.get("author", "")

        for _ in range(8):
            doc.add_paragraph("")

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_title.add_run(title.upper())
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(44, 36, 22)

        if author:
            p_author = doc.add_paragraph()
            p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_a = p_author.add_run(f"\n\n{author}")
            run_a.font.size = Pt(14)
            run_a.font.name = "Times New Roman"
            run_a.font.color.rgb = RGBColor(107, 93, 79)

        chapter_num = 0
        for item in project_data.get("chapters", []):
            chapter_num += 1
            doc.add_page_break()

            content_soup_d = BeautifulSoup(item.get("content", ""), "lxml")
            html_heading_d = content_soup_d.find(["h1", "h2", "h3"])
            chapter_title_d = html_heading_d.get_text(strip=True) if html_heading_d else item.get("title", f"Capítulo {chapter_num}")
            libro_label_d = item.get("libro_title", "").strip()

            for _ in range(3):
                doc.add_paragraph("")

            if libro_label_d:
                p_label = doc.add_paragraph()
                p_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_label = p_label.add_run(libro_label_d.upper())
                run_label.font.size = Pt(9)
                run_label.font.name = "Times New Roman"
                run_label.font.color.rgb = RGBColor(154, 142, 128)
                run_label.font.small_caps = True

            p_ch_title = doc.add_paragraph()
            p_ch_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_ch = p_ch_title.add_run(chapter_title_d)
            run_ch.font.size = Pt(16)
            run_ch.font.bold = True
            run_ch.font.name = "Times New Roman"
            run_ch.font.color.rgb = RGBColor(44, 36, 22)

            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_sep = p_sep.add_run("— ❧ —")
            run_sep.font.size = Pt(10)
            run_sep.font.color.rgb = RGBColor(192, 184, 168)

            doc.add_paragraph("")

            p_running = doc.add_paragraph()
            p_running.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run_running = p_running.add_run(f"— {chapter_title_d} —")
            run_running.font.color.rgb = RGBColor(160, 160, 160)
            run_running.font.size = Pt(9)
            run_running.font.italic = True

            for m in item.get("medias", []):
                if m.get("position") == "before":
                    self._add_media(doc, m)

            soup = clean_html(item.get("content", ""))
            first_heading_skipped_d = False
            for p in soup.find_all(["p", "h1", "h2", "h3"]):
                if p.name in ("h1", "h2", "h3") and not first_heading_skipped_d:
                    first_heading_skipped_d = True
                    continue
                text = p.get_text(strip=True)
                if text:
                    para = doc.add_paragraph(text)
                    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

            for m in item.get("medias", []):
                if m.get("position") in ("after", "inline"):
                    self._add_media(doc, m)

            for note_text in item.get("author_notes", []):
                if note_text:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    run = p.add_run(f"📌 Nota del autor: {note_text}")
                    run.font.color.rgb = RGBColor(140, 123, 107)
                    run.italic = True
                    run.font.size = Pt(10)

        doc.save(output_path)
        return True, f"Borrador DOCX generado en: {output_path}"

    def _add_media(self, doc, m):
        path = m.get("path", "")
        if not os.path.exists(path):
            return
        try:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            run.add_picture(path, width=Inches(4))
            if m.get("caption"):
                cap_para = doc.add_paragraph(m["caption"])
                cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_para.runs[0].italic = True
                cap_para.runs[0].font.size = Pt(9)
                cap_para.runs[0].font.color.rgb = RGBColor(122, 110, 96)
        except Exception:
            pass
